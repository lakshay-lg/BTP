from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import AnalysisResult


NAVY = "12263A"
TEAL = "0E7C7B"
MINT = "DDF4EE"
AMBER = "F2B84B"
PALE = "F4F7F8"
WHITE = "FFFFFF"
RED = "C94C4C"


def _title(ws, title: str, subtitle: str = "") -> None:
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    if subtitle:
        ws.merge_cells("A2:H2")
        ws["A2"] = subtitle
        ws["A2"].font = Font(size=10, color="52636F")
        ws.row_dimensions[2].height = 23


def _table(ws, start_row: int, headers: list[str], rows: list[list]) -> None:
    thin = Side(style="thin", color="D9E1E5")
    for column, header in enumerate(headers, 1):
        cell = ws.cell(start_row, column, header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row_number, row in enumerate(rows, start_row + 1):
        for column, value in enumerate(row, 1):
            cell = ws.cell(row_number, column, value)
            if isinstance(value, str) and value.startswith('='):
                cell.data_type = 's'  # Source/model text must never become an executable formula.
            cell.fill = PatternFill("solid", fgColor=WHITE if row_number % 2 else PALE)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=column in {2, len(headers)})
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
    ws.freeze_panes = f"A{start_row + 1}"
    for column in range(1, len(headers) + 1):
        values = [str(ws.cell(row, column).value or "") for row in range(start_row, min(start_row + len(rows) + 1, start_row + 100))]
        ws.column_dimensions[get_column_letter(column)].width = min(42, max(11, max(map(len, values), default=10) + 2))


def _currency_cells(ws, columns: list[int], start: int, finish: int) -> None:
    for row in range(start, finish + 1):
        for column in columns:
            ws.cell(row, column).number_format = '₹#,##0.00'


def export_xlsx(result: AnalysisResult) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Executive Summary"
    _title(summary, "BuildFlow analysis", "Planner-verifiable draft — not an approved baseline")
    metrics = [
        ["Project", result.project.name],
        ["Data provenance", "Sanitised demonstration rows" if result.project.data_provenance == "sanitised_demo" else "User-uploaded source"],
        ["Detected typology", result.typology.replace("_", " ").title()],
        ["Typology confidence", result.typology_confidence],
        ["BOQ value", result.metrics["boq_total"] if result.metrics["boq_total"] is not None else "Unavailable — source prices missing"],
        ["CPM duration (working days)", result.metrics["duration_working_days"]],
        ["Calculated finish", result.metrics["finish_date"]],
        ["Critical path", " → ".join(result.metrics["critical_path"])],
        ["Classified BOQ rows", result.metrics["classified_percent"] / 100],
    ]
    if result.normalization_review.get('has_overrides'):
        metrics.append(['Input warning', 'Planner-overridden inputs — inspect Input Decisions; these are not verified source facts.'])
    _table(summary, 4, ["Metric", "Value"], metrics)
    summary[9][1].number_format = '₹#,##0.00'
    summary[8][1].number_format = "0%"

    boq = workbook.create_sheet("Priced BOQ")
    _title(boq, "Priced BOQ and classification", "Every AI/retrieval decision is reviewable")
    boq_rows = [
        [item.id, item.description, item.unit, item.quantity, item.source_metadata.get('rate') if item.source_metadata else item.rate, None if item.amount_missing else item.amount, item.work_package, item.track, item.confidence, item.classifier, "; ".join(item.evidence), "; ".join(item.flags)]
        for item in result.items
    ]
    _table(boq, 4, ["ID", "Description", "Unit", "Quantity", "Rate", "Amount", "Work package", "Track", "Confidence", "Classifier", "Evidence", "Flags"], boq_rows)
    _currency_cells(boq, [5, 6], 5, 4 + len(boq_rows))
    for row in range(5, 5 + len(boq_rows)):
        boq.cell(row, 9).number_format = "0%"

    schedule = workbook.create_sheet("CPM Schedule")
    _title(schedule, "Deterministic CPM schedule", "Finish-to-start network generated from the selected typology")
    schedule_rows = [
        [activity.id, activity.name, activity.track, activity.duration_days, ", ".join(activity.predecessors), activity.start_date, activity.finish_date, activity.total_float_days, "Yes" if activity.critical else "No", None if any(i.amount_missing for i in result.items) else activity.cost, ", ".join(activity.boq_item_ids), activity.assumption]
        for activity in result.activities
    ]
    _table(schedule, 4, ["ID", "Activity", "Track", "Duration", "Predecessors", "Start", "Finish", "Float", "Critical", "Allocated cost", "BOQ rows", "Basis / assumption"], schedule_rows)
    _currency_cells(schedule, [10], 5, 4 + len(schedule_rows))
    for row in range(5, 5 + len(schedule_rows)):
        if schedule.cell(row, 9).value == "Yes":
            schedule.cell(row, 9).fill = PatternFill("solid", fgColor=AMBER)

    def add_cash_sheet(name: str, title: str, periods) -> None:
        ws = workbook.create_sheet(name)
        _title(ws, title, "Gross planned value, expenditure and lagged receipts")
        rows = [[p.period, p.gross_work_value, p.site_overhead, p.planned_expenditure, p.certified_receipt, p.retention_withheld, p.net_cash_flow, p.cumulative_work_value, p.cumulative_percent / 100] for p in periods]
        _table(ws, 4, ["Month", "Gross work value", "Site overhead", "Planned expenditure", "Certified receipt", "Retention withheld", "Net cash flow", "Cumulative value", "Cumulative %"], rows)
        _currency_cells(ws, [2, 3, 4, 5, 6, 7, 8], 5, 4 + len(rows))
        for row in range(5, 5 + len(rows)):
            ws.cell(row, 9).number_format = "0.0%"

    add_cash_sheet("Primary Cash Flow", "Primary monthly cash flow", result.cashflow)
    if not result.cashflow:
        workbook['Primary Cash Flow']['A2']='No cash forecast: schedule-only selection or missing source prices.'
    if result.independent_cashflow:
        add_cash_sheet("Independent Phase Curve", "Independent contract-duration cash curve", result.independent_cashflow)

    checks = workbook.create_sheet("Assumptions & Checks")
    _title(checks, "Assumptions and audit findings", "Resolve high-severity items before relying on the output")
    check_rows = [[finding.severity.upper(), finding.code, finding.message, finding.recommendation] for finding in result.findings]
    _table(checks, 4, ["Severity", "Code", "Finding", "Recommended action"], check_rows)
    row = 7 + len(check_rows)
    checks.cell(row, 1, "Model assumptions").font = Font(bold=True, color=WHITE)
    checks.cell(row, 1).fill = PatternFill("solid", fgColor=NAVY)
    for index, assumption in enumerate(result.assumptions, row + 1):
        checks.cell(index, 1, f"A{index - row}")
        checks.cell(index, 2, assumption)
        checks.cell(index, 2).alignment = Alignment(wrap_text=True, vertical="top")
    checks.column_dimensions["A"].width = 16
    checks.column_dimensions["B"].width = 95

    trace = workbook.create_sheet("Agent Trace")
    _title(trace, "Agent orchestration trace", "The agent selects and calls tools; it does not calculate CPM or cash flow itself")
    trace_rows = [[event.step, event.tool, event.status, event.summary, str(event.details)] for event in result.trace]
    _table(trace, 4, ["Step", "Tool", "Status", "Summary", "Details"], trace_rows)

    if result.normalization_review:
        review=workbook.create_sheet('Normalization Review')
        _title(review,'Normalized-source review','Source quotes and human confirmation retained separately from model proposals')
        rows=[[key,str(value)] for key,value in result.normalization_review.items() if key not in {'source_document','decision_audit'}]
        _table(review,4,['Review field','Recorded value'],rows)
        start=len(rows)+7
        _table(review,start,['Item','Source','Original description','Parent/context quotes','Operation','Classification reason'],[
            [i.id,str(i.source_metadata['source']),i.source_metadata['raw_description'],
             str(i.source_metadata['context']),i.source_metadata['operation'],i.source_metadata['classification_reason']]
            for i in result.items])
        if result.normalization_review.get('decision_audit'):
            decisions=workbook.create_sheet('Input Decisions')
            _title(decisions,'Reviewed input decisions','Original evidence retained; planning overrides are not source-verified facts')
            def scalar(value):
                return str(value) if isinstance(value,(dict,list)) else value
            _table(decisions,4,['Item','Field','Source cell','Excel value','AI value','Effective value','Action','Status','Reason','Reviewer','Recorded at'],[
                [d['item_id'],d['field'],d['source_ref'],scalar(d['source_value']),scalar(d['proposed_value']),
                 scalar(d['effective_value']),d['action'],d['status'],d['reason'],d['reviewer'],d['recorded_at']]
                for d in result.normalization_review['decision_audit']])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
