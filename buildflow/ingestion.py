from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import BinaryIO, Iterable

from openpyxl import load_workbook

from .models import BOQItem


HEADER_ALIASES = {
    # A bare "Item" is commonly the serial-number column, so it is deliberately
    # excluded here; accepting it silently shifts every classification to 1, 2, 3…
    "description": {"description", "item description", "particulars", "description of item", "work description"},
    "unit": {"unit", "uom", "units"},
    "quantity": {"quantity", "qty", "estimated quantity", "total qty"},
    "rate": {"rate", "unit rate", "quoted rate", "rate in rs"},
    "amount": {"amount", "total amount", "value", "total", "amount in rs"},
}


def _normalise_header(value: object) -> str:
    text = re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()
    return text


def _number(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    text = re.sub(r"^(?:rs\.?|inr|₹)\s*", "", text, flags=re.I)
    if text in {"", "-", "—", "na", "n/a"}:
        return 0.0
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else 0.0


def _find_columns(row: Iterable[object]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, value in enumerate(row):
        header = _normalise_header(value)
        for field, aliases in HEADER_ALIASES.items():
            if header in aliases and field not in columns:
                columns[field] = index
    return columns


def _items_from_rows(rows: list[list[object]], sheet: str = "") -> list[BOQItem]:
    header_index = -1
    columns: dict[str, int] = {}
    for index, row in enumerate(rows[:40]):
        candidate = _find_columns(row)
        if "description" in candidate and "quantity" in candidate:
            header_index, columns = index, candidate
            break
    if header_index < 0:
        raise ValueError("Could not find a BOQ header row with Description and Quantity columns")

    items: list[BOQItem] = []
    max_col = max(columns.values())
    for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if len(row) <= max_col:
            row = row + [None] * (max_col + 1 - len(row))
        description = str(row[columns["description"]] or "").strip()
        quantity = _number(row[columns["quantity"]])
        unit = str(row[columns.get("unit", -1)] or "").strip() if "unit" in columns else ""
        rate = _number(row[columns["rate"]]) if "rate" in columns else 0.0
        amount = _number(row[columns["amount"]]) if "amount" in columns else 0.0
        if not amount and quantity and rate:
            amount = quantity * rate
        if not description or not quantity:
            continue
        flags: list[str] = []
        if rate <= 0 and amount <= 0:
            flags.append("Unpriced item; excluded from monetary totals")
        elif rate <= 0 and amount > 0 and quantity > 0:
            rate = amount / quantity
            flags.append("Rate derived from amount divided by quantity")
        item_id = f"BOQ-{len(items) + 1:03d}"
        items.append(
            BOQItem(
                id=item_id,
                description=description,
                unit=unit,
                quantity=quantity,
                rate=rate,
                amount=amount,
                source_row=row_index,
                source_sheet=sheet,
                flags=flags,
            )
        )
    if not items:
        raise ValueError("The BOQ contained no rows with both a description and non-zero quantity")
    return items


def read_csv(content: bytes, filename: str = "boq.csv") -> list[BOQItem]:
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    return _items_from_rows(rows, Path(filename).stem)


def read_excel(stream: BinaryIO | io.BytesIO, filename: str = "boq.xlsx") -> list[BOQItem]:
    workbook = load_workbook(stream, data_only=True, read_only=True)
    errors: list[str] = []
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        try:
            items = _items_from_rows(rows, sheet.title)
            if items:
                return items
        except ValueError as exc:
            errors.append(f"{sheet.title}: {exc}")
    raise ValueError("No usable BOQ sheet found. " + "; ".join(errors))


def read_boq(source: str | Path | BinaryIO, filename: str | None = None) -> list[BOQItem]:
    if isinstance(source, (str, Path)):
        path = Path(source)
        filename = filename or path.name
        content = path.read_bytes()
    else:
        filename = filename or getattr(source, "name", "boq.xlsx")
        content = source.read()
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return read_csv(content, filename)
    if suffix in {".xlsx", ".xlsm"}:
        return read_excel(io.BytesIO(content), filename)
    raise ValueError("Supported BOQ formats are .csv, .xlsx, and .xlsm")
