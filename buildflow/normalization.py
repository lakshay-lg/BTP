"""Verify model proposals against uploaded originals; never trust self-checks."""
from __future__ import annotations

import hashlib
from decimal import Decimal

from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

from .normalization_schema import load_document
from .normalization_source import SourceBook, canonical_unit, decimal_value


# Conservative admission list: an available family fallback is NOT evidence of
# a valid rate for a distinct operation (e.g. cleaning hours versus drilling m).
OPERATIONS = {
    'excavation': ('earthwork', {'m3'}),
    'return_fill': ('backfill', {'m3'}),
    'plain_concrete': ('pcc', {'m3'}),
    'reinforced_concrete': ('rcc', {'m3'}),
    'formwork_installation': ('formwork', {'m2','ft2'}),
    'reinforcement_fixing': ('reinforcement', {'kg','t'}),
    'bore_drilling': ('borewell', {'m'}),
    'waterproofing_application': ('waterproofing', {'m2','ft2'}),
}


def review_document(content: bytes, sources: dict[str, bytes]) -> dict:
    document = load_document(content)
    expected = [s['file'] for s in document['source_files']]
    if len(set(expected)) != len(expected) or set(expected) != set(sources):
        raise ValueError('Upload exactly the original files named in source_files; duplicate names are not allowed')
    books = {name:SourceBook(name,data) for name,data in sources.items()}
    checks = []

    def add(code, message, blocks=('extraction',), item_id=None):
        checks.append(dict(code=code,message=message,blocks=list(blocks),item_id=item_id))

    def reference(ref, quote=None):
        if ref['file'] not in books:
            raise ValueError('Evidence references a file that was not uploaded')
        book = books[ref['file']]
        value = book.cell(ref['sheet'],ref['cell'],formulas=True)
        if quote is not None and str(value if value is not None else '') != quote:
            add('SOURCE_TEXT_MISMATCH',f"Quoted text differs from {ref['sheet']}!{ref['cell']}")
        return book.cell(ref['sheet'],ref['cell'])

    inventories = {}
    for inventory in document['source_inventory']:
        key = inventory['file'],inventory['sheet']
        if key in inventories:
            add('DUPLICATE_SHEET','A source sheet is inventoried more than once')
        inventories[key] = inventory
    for filename,book in books.items():
        for sheet in book.formulas.sheetnames:
            populated = book.populated_rows(sheet)
            if not populated:
                continue
            inventory = inventories.get((filename,sheet))
            if not inventory:
                add('MISSING_SHEET',f'Populated source sheet {filename}: {sheet} is not inventoried')
                continue
            if inventory['role'] != 'boq':
                if book.layout(sheet) and book.candidate_rows(sheet):
                    add('BOQ_SHEET_OMITTED',f'{sheet} has BOQ item columns but is excluded; resolve workbook selection')
                continue
            if not book.layout(sheet):
                add('UNSUPPORTED_HEADERS',f'Cannot independently identify Description/Quantity columns in {sheet}')
            listed = inventory['item_rows']+[row['row'] for row in inventory['non_item_rows']]
            if len(listed)!=len(set(listed)) or set(listed)!=populated:
                add('INVENTORY_MISMATCH',f'{sheet}: listed item/non-item rows must uniquely cover every populated row')
            emitted = [i['source']['row'] for i in document['items'] if (i['source']['file'],i['source']['sheet'])==(filename,sheet)]
            candidates = book.candidate_rows(sheet)
            if set(emitted)!=set(inventory['item_rows']) or not candidates <= set(emitted):
                add('MISSING_SOURCE_ITEM',f'{sheet}: source item rows are missing or mislisted; independently found {sorted(candidates-set(emitted))}')
    for filename,sheet in inventories:
        if filename not in books or sheet not in books[filename].formulas.sheetnames:
            add('UNKNOWN_SHEET','Inventory references a sheet absent from the uploads')

    item_ids,source_rows = set(),set()
    for item in document['items']:
        item_id = item['id']
        source = item['source']
        identity = source['file'],source['sheet'],source['row']
        if item_id in item_ids or identity in source_rows:
            add('DUPLICATE_ITEM','Duplicate item ID or source commercial row',item_id=item_id)
        item_ids.add(item_id)
        source_rows.add(identity)
        if source['file'] not in books or source['sheet'] not in books[source['file']].formulas.sheetnames:
            add('UNKNOWN_SOURCE','Item refers to an unavailable source sheet',item_id=item_id)
            continue
        book = books[source['file']]
        layout = book.layout(source['sheet'])
        if not layout:
            continue
        header,columns,price_columns = layout
        if source['row'] not in book.candidate_rows(source['sheet']):
            add('NON_ITEM_ROW',f'{item_id}: row is not independently identifiable as a leaf item',item_id=item_id)
        inventory = inventories.get(identity[:2],{})
        if inventory.get('role')!='boq':
            add('ITEM_OUTSIDE_BOQ','Item does not belong to an included BOQ sheet',item_id=item_id)

        def value_at(cell, allowed=None):
            if not cell:
                return None
            col,row = coordinate_from_string(cell)
            if row!=source['row'] or (allowed is not None and column_index_from_string(col) not in allowed):
                add('WRONG_SOURCE_CELL',f'{item_id}: {cell} is not the matching item/field cell',item_id=item_id)
            return book.cell(source['sheet'],cell)

        try:
            refs=item['description_refs']
            descriptions=[value_at(c,{columns['description']}) for c in refs]
            if len(refs)!=1 or item['raw_description'] != str(descriptions[0] if descriptions[0] is not None else ''):
                add('SOURCE_DESCRIPTION_MISMATCH',f'{item_id}: leaf description differs from source',item_id=item_id)
            for context in item['context']:
                reference(dict(file=source['file'],sheet=source['sheet'],cell=context['cell']),context['text'])
            if item['quantity_ref'] is None:
                add('MISSING_QUANTITY_REFERENCE',f'{item_id}: quantity cell reference required',item_id=item_id)
            source_quantity=decimal_value(value_at(item['quantity_ref'],{columns['quantity']}))
            if source_quantity != (Decimal(str(item['quantity'])) if item['quantity'] is not None else None):
                add('QUANTITY_MISMATCH',f'{item_id}: quantity differs from source',item_id=item_id)
            unit=value_at(item['unit_ref'],{columns['unit']} if 'unit' in columns else set())
            if 'unit' in columns and item['unit_ref'] is None:
                add('MISSING_UNIT_REFERENCE',f'{item_id}: unit column exists but reference is missing',item_id=item_id)
            if (str(unit) if unit is not None else None)!=item['source_unit']:
                add('SOURCE_UNIT_MISMATCH',f'{item_id}: original unit quote differs from source',item_id=item_id)
            if canonical_unit(unit)!=item['unit']:
                add('UNIT_MISMATCH',f'{item_id}: canonical unit differs from source',item_id=item_id)

            price_basis=document['project']['price_basis']
            for field in ('rate','amount'):
                ref=item[field+'_ref']
                available=price_columns[field]
                source_values=[book.values[source['sheet']].cell(source['row'],c).value for c in available]
                if ref is None and any(v is not None for v in source_values) and price_basis not in {'unresolved'}:
                    add('PRICE_REFERENCE_MISSING',f'{item_id}: supplied {field} lacks its source reference',item_id=item_id)
                value_at(ref,available)
                if ref and len(available)>1:
                    bidder=document['project']['selected_bidder']
                    heading=book.column_heading(source['sheet'],column_index_from_string(coordinate_from_string(ref)[0]),header)
                    if price_basis!='selected_bid' or not bidder or bidder.strip()!=heading.strip():
                        if item[field] is not None:
                            add('BIDDER_SELECTION_MISMATCH',f'{item_id}: price column does not match a selected bidder',item_id=item_id)
            rate=decimal_value(value_at(item['rate_ref'],price_columns['rate']))
            if rate != (Decimal(str(item['rate'])) if item['rate'] is not None else None):
                # Nulls intentionally allowed while a populated bidder choice is unresolved.
                if not (price_basis=='unresolved' and item['rate'] is None):
                    add('RATE_MISMATCH',f'{item_id}: rate differs from source',item_id=item_id)
            amount_ref=item['amount_ref']
            raw_amount=book.cell(source['sheet'],amount_ref,formulas=True) if amount_ref else None
            formula=raw_amount if isinstance(raw_amount,str) and raw_amount.startswith('=') else None
            if item['amount_formula']!=formula:
                add('FORMULA_MISMATCH',f'{item_id}: source amount formula is missing or altered',item_id=item_id)
            amount=decimal_value(value_at(amount_ref,price_columns['amount']))
            proposed=Decimal(str(item['amount'])) if item['amount'] is not None else None
            if item['amount_basis']=='quantity_times_rate':
                expected_formulas={f'={item["quantity_ref"]}*{item["rate_ref"]}',f'={item["rate_ref"]}*{item["quantity_ref"]}'}
                ordinary_formula=formula and formula.replace('$','').replace(' ','').upper() in expected_formulas
                if rate is None or source_quantity is None or (formula and not ordinary_formula) or (raw_amount is not None and not formula):
                    add('INVALID_AMOUNT_DERIVATION',f'{item_id}: amount cannot be inferred as quantity × rate',item_id=item_id)
                elif proposed is None or abs(proposed-source_quantity*rate)>Decimal('0.01'):
                    add('AMOUNT_MISMATCH',f'{item_id}: derived amount does not equal quantity × rate',item_id=item_id)
            elif amount!=proposed:
                if not (price_basis=='unresolved' and proposed is None):
                    add('AMOUNT_MISMATCH',f'{item_id}: amount differs from source',item_id=item_id)
            if item['amount_basis']=='source_value' and (proposed is None or amount is None):
                add('MISSING_AMOUNT_VALUE',f'{item_id}: claimed source amount is unavailable',item_id=item_id)
            if item['amount_basis'] in {'missing','unresolved_formula'} and proposed is not None:
                add('INVALID_AMOUNT_BASIS',f'{item_id}: amount basis contradicts its numeric value',item_id=item_id)
            if amount is not None and rate is not None and source_quantity is not None and abs(amount-source_quantity*rate)>Decimal('0.01'):
                add('SOURCE_AMOUNT_CONFLICT',f'{item_id}: source amount differs from quantity × rate; resolve pricing',('cashflow',),item_id)
        except ValueError as exc:
            add('SOURCE_CELL_ERROR',f'{item_id}: {exc}',item_id=item_id)

        capability=OPERATIONS.get(item['operation'])
        if not capability or item['work_package']!=capability[0] or item['unit'] not in capability[1]:
            add('UNSUPPORTED_OPERATION_UNIT',f'{item_id}: no supported production basis for {item["operation"]} / {item["work_package"]} / {item["unit"]}',('schedule',),item_id)
        if item['quantity'] is None or item['quantity']<0:
            add('INVALID_WORKLOAD',f'{item_id}: a nonnegative source quantity is required',('schedule',),item_id)
        if item['quantity_role']!='physical':
            add('QUANTITY_ROLE_REVIEW',f'{item_id}: cost-only/unresolved workload needs a dedicated supported allocation',('schedule',),item_id)
        if item['amount'] is None:
            add('MISSING_PRICE',f'{item_id}: monetary forecast is incomplete',('cashflow',),item_id)
        elif item['amount']<0:
            add('NEGATIVE_PRICE',f'{item_id}: credits/negative amounts need explicit treatment',('cashflow',),item_id)

    issue_ids=[i['id'] for i in document['issues']]
    if len(issue_ids)!=len(set(issue_ids)):
        add('DUPLICATE_ISSUE','Duplicate issue IDs')
    for item in document['items']:
        if not set(item['issue_ids'])<=set(issue_ids):
            add('UNKNOWN_ISSUE_REFERENCE',f'{item["id"]}: issue link does not exist')
    for issue in document['issues']:
        if not set(issue['item_ids'])<=item_ids:
            add('UNKNOWN_ITEM_REFERENCE','Issue references an unknown item')
        for ref in issue['source_refs']:
            try: reference(ref)
            except ValueError as exc: add('SOURCE_REFERENCE_ERROR',str(exc))
        # Remain separate so human-confirmable scope/duration can be resolved at
        # admission without removing any independently detected source error.
        checks.append(dict(code=issue['code'],message=issue['message'],blocks=issue['blocks'],item_id=None,origin='model'))
    for evidence in document['project']['evidence']:
        try: reference(evidence,evidence['text'])
        except ValueError as exc: add('SOURCE_REFERENCE_ERROR',str(exc))
    if document['project']['source_total'] is not None:
        ref=document['project']['source_total_ref']
        try:
            if not ref or decimal_value(reference(ref))!=Decimal(str(document['project']['source_total'])):
                add('SOURCE_TOTAL_MISMATCH','Declared source total is not supported by its cell')
        except ValueError as exc: add('SOURCE_TOTAL_MISMATCH',str(exc))

    items=document['items']
    unpriced=sum(i['amount'] is None for i in items)
    expected_summary=dict(source_leaf_item_count=len(items),emitted_item_count=len(items),
        missing_quantity_count=sum(i['quantity'] is None for i in items),unpriced_item_count=unpriced,
        unclassified_item_count=sum(i['work_package']=='unknown' for i in items),issue_count=len(document['issues']))
    if document['extraction_summary']!=expected_summary:
        add('SUMMARY_MISMATCH','Extraction-summary counts do not match the emitted records')
    priced_total=None if unpriced else float(sum(Decimal(str(i['amount'])) for i in items))
    return dict(schema_version=document['schema_version'],document=document,items=items,checks=checks,
        source_verified=not any('extraction' in c['blocks'] for c in checks),
        schedule_supported=not any(set(c['blocks']) & {'extraction','schedule'} for c in checks),
        prices_complete=not unpriced,priced_total=priced_total,unpriced_item_count=unpriced,
        approval_required=True,source_sha256={name:hashlib.sha256(data).hexdigest() for name,data in sources.items()},
        normalized_sha256=hashlib.sha256(content).hexdigest(),
        limitations=['Source checks do not prove semantic faithfulness. Review reconstructed descriptions and source inventory.',
                     'All classifications are proposals until separately confirmed by a planner.'])
