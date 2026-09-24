"""Read-only workbook evidence for normalization review."""
from __future__ import annotations

import io
import re
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string


def canonical_unit(value):
    value = re.sub(r'[\s.]', '', str(value or '').lower())
    aliases = {
        'm3': ('cum','m3','cubicmetre','cubicmeter'), 'm2': ('sqm','m2','squaremetre','squaremeter'),
        'kg': ('kg','kilogram'), 'm': ('m','meter','metre','rmt','rm'),
        'no': ('no','nos','number','each','ea'), 'h': ('h','hr','hrs','hour','hours'),
        't': ('t','ton','tonne','mt'), 'day': ('day','days'),
        'ls': ('ls','lumpsum'), 'ft2': ('sqft','ft2'), 'point': ('point','points'),
    }
    return next((unit for unit,names in aliases.items() if value in names), 'unknown')


def decimal_value(value):
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        raise ValueError('Boolean cell is not a quantity or price')
    if isinstance(value, (int,float,Decimal)):
        number=Decimal(str(value))
        if not number.is_finite():
            raise ValueError('Non-finite source number')
        return number
    text = re.sub(r'^(?:₹|Rs\.?|INR)\s*', '', str(value).strip(), flags=re.I)
    if not re.fullmatch(r'[+-]?(?:\d+|\d{1,3}(?:,\d{3})+|\d{1,3}(?:,\d{2})*,\d{3})(?:\.\d+)?', text):
        raise ValueError(f'Unresolved numeric cell: {str(value)[:60]}')
    try:
        number = Decimal(text.replace(',', ''))
        if not number.is_finite():
            raise ValueError('Non-finite source number')
        return number
    except InvalidOperation as exc:
        raise ValueError('Unresolved numeric cell') from exc


def header_field(value):
    text = re.sub(r'[^a-z0-9]+',' ',str(value or '').lower()).strip()
    for field,pattern in (
        ('description', r'(?:description(?: of item)?|item description|particulars|work description)'),
        ('quantity', r'(?:qty|quantity|estimated quantity|total qty)(?: .*|)'),
        ('unit', r'(?:unit|units|uom)'),
        ('rate', r'(?:rate|unit rate|quoted rate)(?: rs| in rs| inr| rupees)?'),
        ('amount', r'(?:amount|total amount|value)(?: rs| in rs| inr| rupees)?'),
    ):
        if re.fullmatch(pattern,text):
            return field
    return None


class SourceBook:
    def __init__(self, filename: str, content: bytes):
        if Path(filename).name != filename or '\\' in filename or Path(filename).suffix.lower() not in {'.xlsx','.xlsm'}:
            raise ValueError('Originals must be named XLSX/XLSM uploads, not paths')
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(e.file_size for e in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('Expanded workbook exceeds 50 MB review limit')
            self.formulas = load_workbook(io.BytesIO(content), data_only=False)
            self.values = load_workbook(io.BytesIO(content), data_only=True)
            if len(self.formulas.sheetnames) > 100:
                raise ValueError('Too many workbook sheets')
            if sum(s.max_row*s.max_column for s in self.formulas) > 250000:
                raise ValueError('Workbook review is limited to 250,000 cells')
        except (zipfile.BadZipFile, OSError, KeyError) as exc:
            raise ValueError('Cannot read original workbook') from exc

    def cell(self, sheet, coordinate, formulas=False):
        book = self.formulas if formulas else self.values
        if sheet not in book.sheetnames:
            raise ValueError(f'Unknown source sheet: {sheet}')
        if not coordinate:
            return None
        column,row = coordinate_from_string(coordinate)
        ws = book[sheet]
        if row > ws.max_row or column_index_from_string(column) > ws.max_column:
            raise ValueError(f'Cell outside source sheet: {sheet}!{coordinate}')
        return ws[coordinate].value

    def layout(self, sheet):
        ws = self.formulas[sheet]
        for row in ws.iter_rows(max_row=min(ws.max_row,40)):
            fields = {header_field(c.value): c.column for c in row if header_field(c.value)}
            if {'description','quantity'} <= fields.keys():
                header = row[0].row
                columns = {'rate':set(), 'amount':set()}
                for cells in ws.iter_rows(min_row=header,max_row=min(ws.max_row,header+2)):
                    for c in cells:
                        field = header_field(c.value)
                        if field in columns:
                            columns[field].add(c.column)
                return header,fields,columns
        return None

    def populated_rows(self, sheet):
        return {row[0].row for row in self.formulas[sheet] if any(c.value is not None for c in row)}

    def candidate_rows(self, sheet):
        layout = self.layout(sheet)
        if not layout:
            return set()
        header,fields,_ = layout
        ws = self.formulas[sheet]
        result = set()
        for row in range(header+1,ws.max_row+1):
            description = ws.cell(row,fields['description']).value
            quantity = ws.cell(row,fields['quantity']).value
            unit = ws.cell(row,fields['unit']).value if 'unit' in fields else None
            if not description or header_field(description)=='description':
                continue
            text = str(description).strip().lower()
            if re.match(r'^(?:grand total|total|sub[ -]?total|final amount|discount amount|time\s*line for completion)\b',text):
                continue
            if quantity is not None or (unit is not None and canonical_unit(unit)!='unknown'):
                result.add(row)
        return result

    def column_heading(self, sheet, column, header_row):
        ws = self.formulas[sheet]
        for merged in ws.merged_cells.ranges:
            if merged.min_row <= header_row <= merged.max_row and merged.min_col <= column <= merged.max_col:
                return str(ws.cell(merged.min_row,merged.min_col).value or '')
        return str(ws.cell(header_row,column).value or '')
