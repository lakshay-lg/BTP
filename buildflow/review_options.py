"""Suggestions derived from workbook cells, never from invented AI corrections."""
from collections import Counter

from openpyxl.utils import get_column_letter

from .normalization_source import SourceBook, canonical_unit, decimal_value
from .review_schema import EDIT_FIELDS


def resolution_options(document, sources):
    books = {name:SourceBook(name, data) for name,data in sources.items()}
    identities = Counter((i['source']['file'],i['source']['sheet'],i['source']['row']) for i in document['items'])
    ids = Counter(i['id'] for i in document['items'])
    options = []
    for item in document['items']:
        src = item['source']
        book = books.get(src['file'])
        sheet, row = src['sheet'], src['row']
        identity = (src['file'], sheet, row)
        if not book or sheet not in book.formulas.sheetnames or not book.layout(sheet):
            continue
        if ids[item['id']] != 1 or identities[identity] != 1 or row not in book.candidate_rows(sheet):
            continue
        header, cols, prices = book.layout(sheet)
        cells = {key:f'{get_column_letter(col)}{row}' for key,col in cols.items() if key in {'description','quantity','unit'}}
        for field in ('rate','amount'):
            candidates = prices[field]
            if len(candidates) == 1:
                cells[field] = f'{get_column_letter(next(iter(candidates)))}{row}'
            elif candidates:
                matches = [c for c in candidates if document['project']['price_basis']=='selected_bid'
                    and document['project']['selected_bidder']
                    and book.column_heading(sheet,c,header).strip()==document['project']['selected_bidder'].strip()]
                if len(matches)==1:
                    cells[field] = f'{get_column_letter(matches[0])}{row}'
                elif all(book.cell(sheet,f'{get_column_letter(c)}{row}',formulas=True) is None for c in candidates):
                    ref=item[field+'_ref']
                    if ref in {f'{get_column_letter(c)}{row}' for c in candidates}:
                        cells[field]=ref
            # No price column means genuinely unavailable, not a zero price.
            else:
                cells[field]=None
        evidence={
            'description_refs':[cells['description']],
            'raw_description':str(book.cell(sheet,cells['description']) or ''),
            'quantity_ref':cells['quantity'],
        }
        if 'unit' in cells:
            evidence.update(unit_ref=cells['unit'],source_unit=book.cell(sheet,cells['unit']))
        if 'rate' in cells:
            evidence['rate_ref']=cells['rate']
        if 'amount' in cells:
            evidence['amount_ref']=cells['amount']
            raw=book.cell(sheet,cells['amount'],formulas=True)
            evidence['amount_formula']=raw if isinstance(raw,str) and raw.startswith('=') else None
        try:
            evidence['context']=[dict(cell=c['cell'],text=str(book.cell(sheet,c['cell'],formulas=True) or '')) for c in item['context']]
        except ValueError:
            pass  # Invalid context locations require repair, not acceptance.
        patch={key:value for key,value in evidence.items() if item[key]!=value}
        if patch:
            options.append(dict(item_id=item['id'],field='source_evidence',source_available=True,
                source_value=patch,proposed_value={key:item[key] for key in patch},source_ref=f'{sheet} row {row}',
                actions=['source'],suggestion_label='Source-based evidence repair'))
        for field in EDIT_FIELDS:
            available=False
            value=None
            ref=cells.get(field)
            if field=='unit' and field in cells:
                value=canonical_unit(book.cell(sheet,ref))
                available=True
            elif field in {'quantity','rate','amount'} and field in cells:
                try:
                    value=decimal_value(book.cell(sheet,ref))
                    raw=book.cell(sheet,ref,formulas=True)
                    available=not (isinstance(raw,str) and raw.startswith('=') and value is None)
                    if field=='amount' and not available and cells.get('quantity') and cells.get('rate'):
                        expected={f'={cells["quantity"]}*{cells["rate"]}',f'={cells["rate"]}*{cells["quantity"]}'}
                        qty=decimal_value(book.cell(sheet,cells['quantity']))
                        rate=decimal_value(book.cell(sheet,cells['rate']))
                        if raw.replace('$','').replace(' ','').upper() in expected and qty is not None and rate is not None:
                            value=qty*rate
                            available=True
                    value=float(value) if value is not None else None
                except ValueError:
                    available=False
            options.append(dict(item_id=item['id'],field=field,source_available=available,
                source_value=value,proposed_value=item[field],source_ref=f'{sheet}!{ref}' if ref else None,
                actions=(['source'] if available else [])+['ai','manual'],suggestion_label='Source-based suggestion' if available else 'No direct source suggestion'))
    return options


GUIDANCE = {
    'QUANTITY_MISMATCH':'Compare the source quantity with the AI value. Choose Excel, accept AI, or enter a reviewed quantity. A reason is optional.',
    'UNIT_MISMATCH':'Choose a canonical unit. If changing the physical unit, explicitly review the quantity too; no automatic conversion is performed.',
    'RATE_MISMATCH':'Select the source rate or document the rate assumption used for planning.',
    'AMOUNT_MISMATCH':'Review the source amount and the effective quantity × rate. Choose an amount explicitly.',
    'MISSING_PRICE':'Enter a reviewed price assumption, or use schedule-only output. A reason is optional.',
    'EFFECTIVE_AMOUNT_CONFLICT':'The chosen quantity/rate do not reconcile with the amount. Enter the computed amount or explicitly accept the differing amount. A reason is optional.',
    'SOURCE_AMOUNT_CONFLICT':'Review the differing source amount. An explicit amount decision is required for cash calculations.',
    'UNSUPPORTED_OPERATION_UNIT':'Review the operation, package and unit together. Correct a misclassification only if justified; unsupported work needs a model capability update.',
    'UNIT_QUANTITY_REVIEW':'Explicitly review the quantity when choosing a unit that differs from Excel. No automatic conversion is performed; a reason is optional.',
    'MISSING_SOURCE_ITEM':'Restore the missing source items in the normalized JSON and inventory, then upload again. Do not omit project scope.',
    'INVENTORY_MISMATCH':'Correct item/non-item row coverage in the normalized JSON, then upload again.',
    'UNSUPPORTED_HEADERS':'Correct the original workbook layout or extend the importer so source columns can be verified independently.',
    'DUPLICATE_ITEM':'Remove the duplicate extraction or correct its source identity in the normalized JSON, then reupload.',
    'INVALID_WORKLOAD':'Enter a finite nonnegative quantity, retaining the original source evidence. A reason is optional.',
    'NEGATIVE_PRICE':'Negative/credit amounts need dedicated monetary treatment; do not change their sign merely to pass validation.',
    'QUANTITY_ROLE_REVIEW':'Cost-only or unresolved scope needs a supported allocation; field overrides cannot silently treat it as physical work.',
    'SUMMARY_MISMATCH':'Correct the extraction summary in the original normalized JSON to match its actual records, then reupload.',
}


def explain_checks(checks):
    for check in checks:
        check['guidance']=GUIDANCE.get(check['code'],
            'Inspect the referenced source and use a source-based evidence repair if offered. Otherwise correct the normalized JSON/evidence and reupload; this check cannot be waived.')
        if check.get('origin')=='model':
            check['guidance']='Resolve the model-reported question using source evidence. Correct and reupload the normalized JSON if needed; field overrides do not dismiss model issues.'
    return checks
