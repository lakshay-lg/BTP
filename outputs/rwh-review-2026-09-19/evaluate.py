"""Read-only, reproducible evaluation of the two supplied RWH workbooks.

Run from the repository root:
    PYTHONPATH=. python3 outputs/rwh-review-2026-09-19/evaluate.py
No source workbook or production module is changed. No external LLM is called.
"""
import copy
import hashlib
import json
from dataclasses import asdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from buildflow.agent import ProjectAgent
from buildflow.calendar import working_days_between
from buildflow.ingestion import _items_from_rows, read_boq
from buildflow.models import ProjectConfig
from buildflow.scheduling import build_activities, productivity_for
from buildflow.taxonomy import TAXONOMY, infer_typology

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent.parent
SOURCE = Path('/home/lakshay/Downloads/BOQ')
RAW = SOURCE / 'BOQ RWH (1).xlsx'
COMPANION = SOURCE / 'RWH_BOQ_Schedule_CashFlow (1).xlsx'

# Provisional analyst labels, not supervisor-approved ground truth. Broad family
# labels accept borewell suboperations and cover/frame rows as borewell/manhole;
# that acceptance does NOT validate their production units or activity sequence.
LABELS = {
    10: 'earthwork', 11: 'backfill', 14: 'pcc', 17: 'formwork',
    21: 'reinforcement', 26: 'rcc', 27: 'rcc',
    29: 'filter_media', 30: 'filter_media', 31: 'filter_media',
    33: 'borewell', 34: 'borewell', 36: 'access_metalwork',
    37: 'borewell', 38: 'plumbing', 39: 'manhole', 40: 'manhole', 41: 'borewell',
}
PARENTS = {
    10: [8, 9], 11: [8], 14: [12, 13], 17: [15, 16],
    21: [18, 19, 20], 26: [22, 23, 24, 25], 27: [22, 23, 24, 25],
    **{r: [28] for r in (29, 30, 31)},
    **{r: [32] for r in (33, 34, 36, 37, 38, 39, 40, 41)},
}
PACKAGES = {t.key for t in TAXONOMY}
PHASES = {t.key: t.phase for t in TAXONOMY}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def audit_rows(result, expected):
    return [dict(source_row=i.source_row, description=i.description,
                 quantity=i.quantity, unit=i.unit, amount=i.amount,
                 predicted=i.work_package, expected=expected[i.source_row],
                 supported=expected[i.source_row] in PACKAGES,
                 correct=i.work_package == expected[i.source_row],
                 confidence=i.confidence, evidence=i.evidence,
                 capacity_per_day=productivity_for(i.work_package, i.unit))
            for i in result.items]


def run_case(name, items, expected, mode='hybrid', structures=2, contract=90):
    # Dates/duration are comparison controls from the supplied planning workbook,
    # not actual contract dates or independently observed project duration.
    config = ProjectConfig(name='Rain Water Harvesting, Narsi Village',
                           start_date='2026-04-25', structure_count=structures,
                           contract_duration_days=contract, classifier_mode=mode,
                           payment_lag_months=0, indirect_cost_pct=0,
                           retention_pct=0, use_llm_fallback=False)
    result = ProjectAgent().run(items, config)
    rows = audit_rows(result, expected)
    supported = [r for r in rows if r['supported']]
    summary = dict(case=name, typology=result.typology, **result.metrics,
                   correct_rows=sum(r['correct'] for r in rows),
                   supported_rows=len(supported),
                   correct_supported_rows=sum(r['correct'] for r in supported),
                   high_confidence_errors=[r['source_row'] for r in rows
                                           if not r['correct'] and r['confidence'] >= .68],
                   findings=[f.code for f in result.findings])
    return summary, dict(result=result.to_dict(), row_review=rows)


def main():
    fingerprints = {p.name: sha(p) for p in (RAW, COMPANION)}
    raw_workbook = load_workbook(RAW, data_only=False)
    companion = load_workbook(COMPANION, data_only=False)
    cached = load_workbook(COMPANION, data_only=True)
    sheet = raw_workbook['CS for RWH']
    raw_items = read_boq(RAW)
    assert [i.source_row for i in raw_items] == list(LABELS)
    summaries, details = [], {}
    for mode in ('rules', 'retrieval', 'hybrid'):
        name = 'raw_' + mode
        s, d = run_case(name, raw_items, LABELS, mode)
        summaries.append(s)
        details[name] = d
    for count in (1,):
        name = 'raw_hybrid_one_structure'
        s, d = run_case(name, raw_items, LABELS, structures=count)
        summaries.append(s)
        details[name] = d

    enriched = copy.deepcopy(raw_items)
    for item in enriched:
        item.description = ' '.join(str(sheet.cell(r, 2).value) for r in PARENTS[item.source_row]) + ' ' + item.description
    s, d = run_case('raw_parent_context_hybrid', enriched, LABELS)
    summaries.append(s)
    details[s['case']] = d

    try:
        read_boq(COMPANION)
        companion_import = 'Unexpected success'
    except ValueError as exc:
        companion_import = str(exc)

    # Controlled import adaptation in memory only; the production parser remains
    # unchanged and is evaluated as failing on this file above.
    priced_sheet = cached['Priced BOQ']
    adapted_rows = [list(r) for r in priced_sheet.iter_rows(values_only=True)]
    adapted_rows[3][3:6] = ['Quantity', 'Rate', 'Amount']
    adapted = _items_from_rows(adapted_rows, 'Priced BOQ')
    companion_labels = {r: LABELS[source_row] for r, source_row in zip(range(5, 23), LABELS)}
    for mode in ('rules', 'retrieval', 'hybrid'):
        name = 'companion_header_adapter_' + mode
        s, d = run_case(name, adapted, companion_labels, mode)
        summaries.append(s)
        details[name] = d

    # Oracle-label diagnostic: replace only supported family labels, mark the
    # four genuinely missing categories unknown. Does not claim a deployable fix.
    oracle = copy.deepcopy(raw_items)
    for item in oracle:
        expected = LABELS[item.source_row]
        item.work_package = expected if expected in PACKAGES else 'unknown'
        item.phase = PHASES.get(item.work_package, 'unclassified')
        item.track = 'borewell' if expected == 'borewell' else 'general'
    config = ProjectConfig(name='RWH diagnostic', start_date='2026-04-25', structure_count=2)
    oracle_activities = build_activities(oracle, 'rwh', config)
    known_only_activities = build_activities([i for i in oracle if i.work_package != 'unknown'], 'rwh', config)

    priced_amounts = [priced_sheet.cell(r, 4).value * priced_sheet.cell(r, 5).value for r in range(5, 23)]
    cash = companion['Cash Flow']
    monthly = [sum(cash.cell(r, c).value for r in range(2, 12)) for c in range(2, 7)]
    cash_total = sum(monthly)
    total = sum(priced_amounts)
    ratios = [dict(raw_row=i.source_row, companion_row=j, raw_quantity=i.quantity,
                   companion_quantity=priced_sheet.cell(j, 4).value,
                   ratio=priced_sheet.cell(j, 4).value / i.quantity)
              for i, j in zip(raw_items, range(5, 23))]
    schedule = companion['Schedule']
    calendar_checks = []
    for row in range(3, 27):
        start, end, stated = [schedule.cell(row, c).value for c in (3, 4, 5)]
        counted = len(working_days_between(start.date(), end.date(), 6))
        calendar_checks.append(dict(row=row, name=schedule.cell(row, 1).value,
                                    stated=stated, counted=counted, matches=stated == counted))
    formula_cache_missing = []
    for s in companion:
        for cells in s:
            for cell in cells:
                if cell.data_type == 'f' and cached[s.title][cell.coordinate].value is None:
                    formula_cache_missing.append(f'{s.title}!{cell.coordinate}')
    oracle_summary = {
        'all_rows_days': max(a.finish_day for a in oracle_activities),
        'known_rows_only_days': max(a.finish_day for a in known_only_activities),
        'known_rows_only_warning': 'Excludes four rows of scope; diagnostic, not a valid complete schedule',
        'all_rows_activities': [asdict(a) for a in oracle_activities],
    }
    # Demonstrate that an absent contract duration falls back to CPM, violating
    # the intended independence of the cash project.
    s, d = run_case('companion_no_contract_duration', adapted, companion_labels, contract=None)
    summaries.append(s)
    details[s['case']] = d

    checks = dict(
        source_fingerprints=fingerprints,
        source_units_and_quantities_preserved=all(i.quantity == sheet.cell(i.source_row,4).value for i in raw_items),
        source_bid_price_cells_populated=sum(sheet.cell(r,c).value is not None for r in LABELS for c in range(5,11)),
        source_contract_duration_cells=[sheet.cell(50,c).value for c in range(4,11)],
        companion_import_error=companion_import,
        companion_boq_total_recomputed=total, companion_cash_monthly=monthly,
        companion_cash_total_recomputed=cash_total,
        companion_cash_minus_boq=cash_total-total,
        companion_cash_shortfall_percent=100*(total-cash_total)/total,
        companion_grand_cumulative_formula=cash['G14'].value,
        companion_grand_cumulative_recomputed=2*cash_total,
        quantity_ratios=ratios,
        reference_working_days=len(working_days_between(date(2026,4,25),date(2026,8,7))),
        reference_calendar_days_inclusive=105,
        schedule_row_calendar_checks=calendar_checks,
        missing_formula_caches=formula_cache_missing,
        typology_with_neutral_name=infer_typology(raw_items, 'Untitled project'),
        oracle_label_diagnostic=oracle_summary,
    )
    code_hashes = {p.name: sha(p) for p in (ROOT/'buildflow').glob('*.py')}
    assert fingerprints == {p.name: sha(p) for p in (RAW, COMPANION)}, 'Source changed during analysis'
    payload = dict(method='Offline diagnostic, provisional analyst labels, no production changes',
                   code_sha256=code_hashes, checks=checks, runs=summaries, details=details)
    (OUT/'results.json').write_text(json.dumps(payload, indent=2), encoding='utf-8')
    print(json.dumps(dict(runs=summaries, checks={k:v for k,v in checks.items() if k not in ('oracle_label_diagnostic',)}), indent=2))


if __name__ == '__main__':
    main()
