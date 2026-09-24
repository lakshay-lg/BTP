"""Compare the generated two-tank schedule with the original workbook's Schedule and Cash Flow sheets.

usage: python3 scripts/compare_original.py <original.xlsx> [generated_schedule.json]
"""
import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl

MAP = {"Mobilization & Site Setup": ["MOB"], "Excavation": ["A01"], "PCC Base": ["A02"],
       "Formwork & Reinforcement": ["A03", "A04"], "RCC Pour (footing/wall/slab)": ["A05", "A06"],
       "Curing & Deshuttering": ["CUR"], "Backfill": ["A07"],
       "Gravel Filter Bed (boulder/gravel/sand)": ["A08", "A09", "A10"],
       "Borewell Drilling 0-30m": ["A11"], "Screen Pipe & Vent Installation": ["A12", "A13", "A16"],
       "Manholes & Foot Rests": ["A15", "A17", "A18"], "Bore Cleaning": ["A14"],
       "Joint Testing & Handover": ["TEST"]}


def wd(a, b):  # working days between two dates inclusive, Sunday off
    return sum(1 for i in range((b - a).days + 1) if (a + dt.timedelta(i)).weekday() != 6)


def main(xlsx, gen="data/rwh_two_tank_schedule.json"):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    rows = [r for r in wb["Schedule"].iter_rows(values_only=True)
            if isinstance(r[2], dt.datetime) and isinstance(r[3], dt.datetime) and r[1]]
    d = json.loads(Path(gen).read_text())
    by = {a["activity_id"]: a for a in d["activities"]}
    p = lambda s: dt.date.fromisoformat(s)
    o0 = min(r[2] for r in rows).date()
    print(f"{'Activity':46} {'ORIGINAL start-end (days)':27} {'GENERATED start-end (days)':27} end diff")
    for r in rows:
        name = r[0]; tank = None
        if name.startswith("T1: ") or name.startswith("T2: "):
            tank, base = name[1], name[4:]
        else:
            base = name
        ids = [(f"T{tank}_" + i) if tank else i for i in MAP[base]]
        g = [by[i] for i in ids]
        gs, ge = min(p(x["start_date"]) for x in g), max(p(x["end_date"]) for x in g)
        os_, oe = r[2].date(), r[3].date()
        print(f"{name[:46]:46} {os_:%d %b}-{oe:%d %b} ({r[4]:>2}d)".ljust(75)
              + f"{gs:%d %b}-{ge:%d %b} ({wd(gs, ge):>2}d)".ljust(28) + f"{(ge - oe).days:+d} d")
    oe_all = max(r[3] for r in rows).date()
    print(f"\nProject: original {o0:%d %b %Y} -> {oe_all:%d %b %Y} ({wd(o0, oe_all)} working days) | "
          f"generated {d['project_start_date']} -> {d['project_end_date']} ({d['project_duration_days']} working days)")
    print("Generated critical path:", " -> ".join(d["critical_path"]))

    # S-curve at the end of each 30-day block from project start
    cf = [r for r in wb["Cash Flow"].iter_rows(min_row=2, values_only=True)
          if r[0] and isinstance(r[6], (int, float)) and "TOTAL" not in str(r[0]) and r[0] != "CUMULATIVE"]
    ototal = sum(r[6] for r in cf)
    ocum, run = [], 0
    for k in range(1, 6):
        run += sum(r[k] for r in cf); ocum.append(100 * run / ototal)
    # original S-curve derived from its own Schedule dates: BOQ amounts (per tank = half of the priced
    # BOQ, which prices the bore-well rows once) spread over each schedule row's working days
    boq = {r[0]: r[3] * r[4] for r in wb["Priced BOQ"].iter_rows(values_only=True)
           if r[0] and isinstance(r[3], (int, float)) and isinstance(r[4], (int, float))}
    rowmap = {"Excavation": ["A.1a"], "PCC Base": ["B.1a"], "Formwork & Reinforcement": ["C.1a", "D.1a"],
              "RCC Pour (footing/wall/slab)": ["E.1", "E.1a"], "Backfill": ["A.2"],
              "Gravel Filter Bed (boulder/gravel/sand)": ["F.1", "F.2", "F.3"], "Borewell Drilling 0-30m": ["G.1"],
              "Screen Pipe & Vent Installation": ["G.2", "G.4", "G.5"], "Manholes & Foot Rests": ["G.3", "G.6", "G.7"],
              "Bore Cleaning": ["G.8"]}
    otasks = []
    for r in rows:
        base = r[0][4:] if r[0][:4] in ("T1: ", "T2: ") else r[0]
        if base in rowmap:
            otasks.append((r[2].date(), r[3].date(), sum(boq[k] for k in rowmap[base]) / 2))
    ototal_sched = sum(t[2] for t in otasks)
    def orig_cum(day):
        cutoff = o0 + dt.timedelta(day); s = 0.0
        for a, b, cost in otasks:
            days = [x for x in (a + dt.timedelta(i) for i in range((b - a).days + 1)) if x.weekday() != 6] or [a]
            s += cost * sum(1 for x in days if x < cutoff) / len(days)
        return 100 * s / ototal_sched
    total = sum(a["cost"] or 0 for a in d["activities"])
    def gen_cum(day):
        cutoff = o0 + dt.timedelta(day)
        s = 0.0
        for a in d["activities"]:
            if not a["cost"]:
                continue
            days = [x for x in (p(a["start_date"]) + dt.timedelta(i) for i in range((p(a["end_date"]) - p(a["start_date"])).days + 1))
                    if x.weekday() != 6] or [p(a["start_date"])]
            s += a["cost"] * sum(1 for x in days if x < cutoff) / len(days)
        return 100 * s / total
    print("\nCumulative % of work value at the end of each 30-day block from the start")
    print("block   original(from its schedule)   original(cash-flow sheet)   generated")
    for k in range(5):
        print(f"{k + 1:>5}   {orig_cum(30 * (k + 1)):>15.1f}%   {ocum[k]:>21.1f}%   {gen_cum(30 * (k + 1)):>8.1f}%")
    print(f"\nTotals: original priced BOQ (qty x rate) {sum(boq.values()):,.0f} | original schedule-basis (BOQ items only) {ototal_sched:,.0f} | original cash-flow sheet {ototal:,.0f} | generated (DAR rates, 2 tanks, incl. GST/CPOH) {total:,.0f}")


if __name__ == "__main__":
    main(*sys.argv[1:])
