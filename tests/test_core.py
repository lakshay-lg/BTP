from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from buildflow.agent import ProjectAgent
from buildflow.benchmark import run_benchmark
from buildflow.exporter import export_xlsx
from buildflow.ingestion import read_boq
from buildflow.models import ProjectConfig
from buildflow.scheduling import apply_duration_overrides, topological_order


ROOT = Path(__file__).resolve().parent.parent


def config(typology: str, days: int, structures: int = 1) -> ProjectConfig:
    return ProjectConfig(
        name=f"{typology} test",
        typology=typology,
        start_date="2026-04-25",
        contract_duration_days=days,
        structure_count=structures,
        indirect_cost_pct=5,
        retention_pct=5,
    )


def test_csv_uses_description_not_serial_number_column():
    items = read_boq(ROOT / "data/samples/stp_boq.csv")
    assert items[0].description.startswith("Earth work excavation")
    assert items[0].description != "1"


def test_excel_header_detection_and_derived_amount():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Tender title"])
    sheet.append(["S.No", "Description of Item", "UOM", "Qty", "Rate"])
    sheet.append([1, "Plain cement concrete below foundation", "cum", 10, 5000])
    stream = BytesIO()
    workbook.save(stream)
    items = read_boq(BytesIO(stream.getvalue()), "test.xlsx")
    assert len(items) == 1
    assert items[0].amount == 50_000


@pytest.mark.parametrize(
    ("sample", "typology", "days", "structures"),
    [("stp", "stp_tank", 150, 1), ("mep", "linear_mep", 130, 1), ("rwh", "rwh", 90, 2)],
)
def test_end_to_end_typologies_reconcile_and_have_valid_cpm(sample, typology, days, structures):
    result = ProjectAgent().run(read_boq(ROOT / f"data/samples/{sample}_boq.csv"), config(typology, days, structures))
    assert result.metrics["classified_percent"] >= 90
    assert result.metrics["cost_reconciliation_delta"] == pytest.approx(0, abs=1)
    assert len(topological_order(result.activities)) == len(result.activities)
    assert any(activity.critical for activity in result.activities)
    assert 10 < result.metrics["duration_working_days"] < 500
    assert result.cashflow[-1].cumulative_percent == pytest.approx(100)
    assert result.independent_cashflow[-1].cumulative_percent == pytest.approx(100)


def test_extra_lift_is_cost_only_not_double_counted_volume():
    result = ProjectAgent().run(read_boq(ROOT / "data/samples/stp_boq.csv"), config("stp_tank", 150))
    lift = next(item for item in result.items if "Extra lift" in item.description)
    excavation = next(activity for activity in result.activities if activity.id == "EXC")
    assert any("Cost-only" in flag for flag in lift.flags)
    assert excavation.duration_days == 44  # 5200 / 120 rounded up; lift quantity excluded


def test_mep_tracks_run_in_parallel():
    result = ProjectAgent().run(read_boq(ROOT / "data/samples/mep_boq.csv"), config("linear_mep", 130))
    by_id = {activity.id: activity for activity in result.activities}
    assert by_id["SEX"].start_day == by_id["TEX"].start_day == by_id["PEX"].start_day


def test_rwh_borewell_is_independent_specialist_track():
    result = ProjectAgent().run(read_boq(ROOT / "data/samples/rwh_boq.csv"), config("rwh", 90, 2))
    by_id = {activity.id: activity for activity in result.activities}
    assert by_id["BOR"].predecessors == ["MOB"]
    assert by_id["T2E"].predecessors == ["T1P"]


def test_independent_cashflow_does_not_change_when_cpm_is_replanned():
    agent = ProjectAgent()
    result = agent.run(read_boq(ROOT / "data/samples/rwh_boq.csv"), config("rwh", 90, 2))
    before = [(period.period, period.gross_work_value) for period in result.independent_cashflow]
    updated = agent.replan(result, [{"id": "BOR", "duration_days": 25}])
    after = [(period.period, period.gross_work_value) for period in updated.independent_cashflow]
    assert before == after
    assert updated.metrics["duration_working_days"] >= result.metrics["duration_working_days"]


def test_cycle_override_is_rejected():
    result = ProjectAgent().run(read_boq(ROOT / "data/samples/rwh_boq.csv"), config("rwh", 90, 2))
    with pytest.raises(ValueError, match="cycle"):
        apply_duration_overrides(result.activities, [{"id": "MOB", "predecessors": ["HND"]}], result.project.start_date, 6)


def test_export_contains_audit_sheets():
    result = ProjectAgent().run(read_boq(ROOT / "data/samples/rwh_boq.csv"), config("rwh", 90, 2))
    workbook = load_workbook(BytesIO(export_xlsx(result)), read_only=True)
    assert {"Priced BOQ", "CPM Schedule", "Primary Cash Flow", "Independent Phase Curve", "Agent Trace"}.issubset(workbook.sheetnames)


def test_offline_ablation_benchmark_runs():
    report = run_benchmark(ROOT / "data/benchmark_labels.csv")
    assert {entry["mode"] for entry in report} == {"rules", "retrieval", "hybrid"}
    hybrid = next(entry for entry in report if entry["mode"] == "hybrid")
    assert hybrid["accuracy"] >= 0.85
