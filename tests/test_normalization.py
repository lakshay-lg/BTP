import copy
import io
import json

import pytest
from openpyxl import Workbook
from openpyxl import load_workbook

from buildflow.web import create_app


def fixture_pair(priced=True):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'BOQ'
    sheet.append(['RWH example'])
    sheet.append(['No', 'Description', 'Unit', 'Qty', 'Rate', 'Amount'])
    sheet.append([1, 'Filling available excavated earth', 'Cum', 12, 10 if priced else None, 120 if priced else None])
    sheet.append([2, 'PCC blinding', 'Cum', 5, 20 if priced else None, 100 if priced else None])
    output = io.BytesIO()
    workbook.save(output)
    items = []
    for index, (desc, qty, rate, amount, package, operation) in enumerate([
        ('Filling available excavated earth', 12, 10, 120, 'backfill', 'return_fill'),
        ('PCC blinding', 5, 20, 100, 'pcc', 'plain_concrete'),
    ], start=3):
        items.append(dict(id=f'I{index}', source=dict(file='original.xlsx', sheet='BOQ', row=index),
                          source_item_number=str(index-2), section=None, raw_description=desc,
                          description_refs=[f'B{index}'], context=[], normalized_description=desc,
                          quantity=qty, quantity_ref=f'D{index}', source_unit='Cum', unit_ref=f'C{index}',
                          unit='m3', quantity_basis='unresolved', quantity_role='physical',
                          rate=rate if priced else None, rate_ref=f'E{index}',
                          amount=amount if priced else None, amount_ref=f'F{index}', amount_formula=None,
                          amount_basis='source_value' if priced else 'missing', work_package=package,
                          operation=operation, classification_reason='Source operation',
                          review_status='proposed', issue_ids=[]))
    doc = dict(schema_version='buildflow.normalized-boq.v0.1', complete=True,
               source_files=[dict(file='original.xlsx', role='original_boq')],
               project=dict(name=None, typology_proposal='rwh', structure_count=None, quantity_basis='unresolved',
                            contract_duration=None, contract_duration_basis=None, start_date=None,
                            price_basis='single_price_column' if priced else 'unpriced', selected_bidder=None,
                            currency='INR', source_total=None, source_total_ref=None, evidence=[]),
               source_inventory=[dict(file='original.xlsx', sheet='BOQ', role='boq', item_rows=[3,4],
                                      non_item_rows=[dict(row=1,role='title',reason='Title'),dict(row=2,role='header',reason='Column labels')])],
               items=items, issues=[], extraction_summary=dict(source_leaf_item_count=2, emitted_item_count=2,
                   missing_quantity_count=0, unpriced_item_count=0 if priced else 2, unclassified_item_count=0,issue_count=0))
    return doc, output.getvalue()


def post_review(doc, source):
    return create_app().test_client().post('/api/normalization/review', data={
        'normalized': (io.BytesIO(json.dumps(doc).encode()), 'normalized_boq.json'),
        'originals': (io.BytesIO(source), 'original.xlsx'),
    })


def test_review_preserves_prices_and_separates_human_approval():
    doc, source = fixture_pair()
    response = post_review(doc, source)
    assert response.status_code == 200
    result = response.get_json()
    assert result['source_verified'] is True
    assert result['priced_total'] == 220
    assert result['items'][0]['work_package'] == 'backfill'
    assert result['approval_required'] is True
    assert len(result['source_sha256']['original.xlsx']) == 64


@pytest.mark.parametrize('change', ['quantity', 'quote', 'omission', 'duplicate', 'wrong_cell'])
def test_source_disagreement_cannot_pass_review(change):
    doc, source = fixture_pair()
    if change == 'quantity':
        doc['items'][0]['quantity'] = 999
    elif change == 'quote':
        doc['items'][0]['context'] = [dict(cell='B1',text='Made-up approved specification')]
    elif change == 'omission':
        doc['items'].pop()
        doc['source_inventory'][0]['item_rows'] = [3]
        doc['source_inventory'][0]['non_item_rows'].append(dict(row=4, role='note', reason='Pretend not an item'))
        doc['extraction_summary'].update(source_leaf_item_count=1,emitted_item_count=1)
    elif change == 'duplicate':
        doc['items'][1] = copy.deepcopy(doc['items'][0])
    else:
        doc['items'][0]['quantity_ref'] = 'E3'
        doc['items'][0]['quantity'] = 10
    response = post_review(doc, source)
    assert response.status_code == 200
    result = response.get_json()
    assert result['source_verified'] is False
    assert any('extraction' in issue['blocks'] for issue in result['checks'])


def test_missing_prices_stay_null():
    doc, source = fixture_pair(priced=False)
    result = post_review(doc, source).get_json()
    assert result['source_verified'] is True
    assert result['priced_total'] is None
    assert result['items'][0]['amount'] is None
    assert result['unpriced_item_count'] == 2


def test_incompatible_workload_units_are_blocked_even_with_valid_sources():
    doc, source = fixture_pair()
    doc['items'][0].update(work_package='testing',operation='system_testing')
    response = post_review(doc, source)
    assert response.status_code == 200
    result = response.get_json()
    assert result['source_verified'] is True
    assert any(x['code']=='UNSUPPORTED_OPERATION_UNIT' for x in result['checks'])


@pytest.mark.parametrize('raw', [b'{"complete":true,"complete":false}', b'{"a":NaN}', b'[]', b'null'])
def test_invalid_json_is_a_client_error(raw):
    _, source = fixture_pair()
    response = create_app().test_client().post('/api/normalization/review', data={
        'normalized': (io.BytesIO(raw),'normalized.json'),
        'originals': (io.BytesIO(source),'original.xlsx'),
    })
    assert response.status_code == 400


@pytest.mark.parametrize('mutation', ['extra', 'false_complete', 'bool_quantity', 'bad_issue_link'])
def test_strict_contract_rejects_invalid_claims(mutation):
    doc, source = fixture_pair()
    if mutation=='extra': doc['items'][0]['approved'] = True
    if mutation=='false_complete': doc['complete'] = False
    if mutation=='bool_quantity': doc['items'][0]['quantity'] = True
    if mutation=='bad_issue_link': doc['items'][0]['issue_ids'] = ['missing']
    response = post_review(doc, source)
    assert response.status_code == 400 or (response.status_code == 200 and not response.get_json()['source_verified'])


def analysis_form(doc, source, **overrides):
    form = dict(normalized=(io.BytesIO(json.dumps(doc).encode()),'normalized.json'),
                originals=(io.BytesIO(source),'original.xlsx'), reviewer='Test planner',
                review_confirmed='true', scope_confirmed='true', quantity_basis='project_total',
                typology='rwh', start_date='2026-04-25', structure_count='1',
                cashflow_mode='schedule', contract_duration_days='90', contract_confirmed='true')
    form.update(overrides)
    return form


def test_reviewed_analysis_keeps_labels_and_records_source_hash():
    doc, source = fixture_pair()
    response = create_app().test_client().post('/api/normalization/analyze', data=analysis_form(doc,source))
    assert response.status_code == 200
    result = response.get_json()
    assert result['items'][0]['work_package']=='backfill'  # old rules incorrectly choose earthwork
    assert result['items'][0]['classifier']=='planner_reviewed'
    assert result['metrics']['boq_total']==220
    assert result['normalization_review']['reviewer']=='Test planner'


@pytest.mark.parametrize('override', [dict(review_confirmed='false'),dict(scope_confirmed='false'),dict(reviewer=''),dict(quantity_basis='unresolved')])
def test_human_review_is_required(override):
    doc, source = fixture_pair()
    response = create_app().test_client().post('/api/normalization/analyze',data=analysis_form(doc,source,**override))
    assert response.status_code == 422


def test_approval_cannot_override_changed_source_quantity():
    doc, source = fixture_pair()
    doc['items'][0]['quantity']=999
    response=create_app().test_client().post('/api/normalization/analyze',data=analysis_form(doc,source))
    assert response.status_code==422


def test_missing_prices_allow_schedule_but_not_cashflow_or_fake_zero():
    doc, source = fixture_pair(False)
    client=create_app().test_client()
    response=client.post('/api/normalization/analyze',data=analysis_form(doc,source))
    assert response.status_code==200
    result=response.get_json()
    assert result['metrics']['boq_total'] is None
    assert result['cashflow']==[]
    assert result['items'][0]['amount_missing'] is True
    response=client.post('/api/normalization/analyze',data=analysis_form(doc,source,cashflow_mode='compare'))
    assert response.status_code==422


def test_independent_curve_cannot_borrow_cpm_duration():
    doc, source = fixture_pair()
    response=create_app().test_client().post('/api/normalization/analyze',data=analysis_form(doc,source,cashflow_mode='compare',contract_duration_days='',contract_confirmed='false'))
    assert response.status_code==422


def change_source(source, cell, value):
    workbook=load_workbook(io.BytesIO(source))
    workbook['BOQ'][cell]=value
    output=io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_zero_amount_is_not_missing_or_replaced_by_product():
    doc,source=fixture_pair()
    source=change_source(source,'F3',0)
    doc['items'][0]['amount']=0
    result=post_review(doc,source).get_json()
    assert result['source_verified'] is True
    assert result['unpriced_item_count']==0
    assert result['items'][0]['amount']==0
    assert result['priced_total']==100
    assert any(c['code']=='SOURCE_AMOUNT_CONFLICT' for c in result['checks'])


@pytest.mark.parametrize('formula,verified',[('=D3*E3',True),('=D3*E3*1.18',False)])
def test_only_plain_source_products_can_be_derived_without_cache(formula,verified):
    doc,source=fixture_pair()
    source=change_source(source,'F3',formula)
    doc['items'][0].update(amount_basis='quantity_times_rate',amount_formula=formula)
    result=post_review(doc,source).get_json()
    assert result['source_verified'] is verified


def test_source_file_name_cannot_be_a_path():
    doc,source=fixture_pair()
    response=create_app().test_client().post('/api/normalization/review',data={
        'normalized':(io.BytesIO(json.dumps(doc).encode()),'normalized.json'),
        'originals':(io.BytesIO(source),'../../original.xlsx'),
    })
    assert response.status_code==400


def test_reviewed_export_and_replan_keep_missing_prices_unavailable():
    doc,source=fixture_pair(False)
    client=create_app().test_client()
    response=client.post('/api/normalization/analyze',data=analysis_form(doc,source))
    result=response.get_json()
    replan=client.post(f'/api/jobs/{result["job_id"]}/replan',json={'overrides':[{'id':'MOB','duration_days':3}]}).get_json()
    assert replan['metrics']['boq_total'] is None
    assert replan['cashflow']==[]
    assert replan['normalization_review']['reviewer']=='Test planner'
    export=client.get(f'/api/jobs/{result["job_id"]}/export.xlsx')
    book=load_workbook(io.BytesIO(export.data))
    assert book['Priced BOQ']['F5'].value is None
    assert 'Normalization Review' in book.sheetnames
    assert 'No cash forecast' in book['Primary Cash Flow']['A2'].value


def test_reviewed_text_is_not_exported_as_a_formula():
    doc,source=fixture_pair()
    doc['items'][0]['normalized_description']='=1+2'
    client=create_app().test_client()
    result=client.post('/api/normalization/analyze',data=analysis_form(doc,source)).get_json()
    book=load_workbook(io.BytesIO(client.get(f'/api/jobs/{result["job_id"]}/export.xlsx').data))
    assert book['Priced BOQ']['B5'].value=='=1+2'
    assert book['Priced BOQ']['B5'].data_type=='s'


def test_prompt_schema_and_page_are_accessible():
    client=create_app().test_client()
    assert client.get('/normalize').status_code==200
    assert client.get('/api/normalization/prompt').status_code==200
    schema=client.get('/api/normalization/schema').get_json()
    assert schema['properties']['schema_version']['const']=='buildflow.normalized-boq.v0.1'


def test_broken_workbook_returns_client_error():
    doc,_=fixture_pair()
    response=post_review(doc,b'not an xlsx')
    assert response.status_code==400


def test_tiny_numeric_source_quantity_is_preserved():
    doc,source=fixture_pair(False)
    source=change_source(source,'D3',0.00001)
    doc['items'][0]['quantity']=0.00001
    assert post_review(doc,source).get_json()['source_verified'] is True
