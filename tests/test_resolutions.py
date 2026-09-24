import copy
import hashlib
import io
import json

import pytest
from openpyxl import load_workbook

from buildflow.web import create_app
from test_normalization import fixture_pair, analysis_form, change_source


def envelope(doc, source, decisions=()):
    return dict(schema_version='buildflow.reviewed-boq.v0.1', normalized_document=doc,
                source_sha256={'original.xlsx':hashlib.sha256(source).hexdigest()},
                decisions=list(decisions), review_settings={}, review_state='draft')


def decision(field='quantity', value=999, action='ai', item='I3', reason='Verified revised site scope'):
    return dict(item_id=item, field=field, action=action, value=value, reason=reason,
                reviewer='Test planner', recorded_at='')


def review(doc, source, client=None):
    client=client or create_app().test_client()
    return client.post('/api/normalization/review', data={
        'normalized':(io.BytesIO(json.dumps(doc).encode()),'reviewed.json'),
        'originals':(io.BytesIO(source),'original.xlsx'),
    })


def test_source_suggestion_and_guidance_are_returned_for_quantity_mismatch():
    doc,source=fixture_pair()
    doc['items'][0]['quantity']=999
    result=review(doc,source).get_json()
    option=next(o for o in result['resolution_options'] if o['item_id']=='I3' and o['field']=='quantity')
    assert option['source_value']==12
    assert option['proposed_value']==999
    assert option['source_ref']=='BOQ!D3'
    assert all(c['guidance'] for c in result['checks'])


@pytest.mark.parametrize('action,value,override', [('source',12,False),('ai',999,True),('manual',24,True)])
def test_decisions_drive_effective_inputs_without_changing_evidence(action,value,override):
    doc,source=fixture_pair()
    doc['items'][0]['quantity']=999
    saved=envelope(doc,source,[decision(value=value,action=action)])
    response=review(saved,source)
    assert response.status_code==200,response.get_json()
    result=response.get_json()
    assert result['document']['items'][0]['quantity']==999
    assert result['items'][0]['quantity']==value
    assert result['has_overrides'] is override
    assert result['source_verified'] is False
    assert result['admission_ready'] is True
    assert result['decision_audit'][0]['source_value']==12
    assert result['reviewed_document']['decisions'][0]['recorded_at']
    client=create_app().test_client()
    run=client.post('/api/normalization/analyze',data=analysis_form(saved,source))
    assert run.status_code==200,run.get_json()
    assert run.get_json()['items'][0]['quantity']==value
    assert run.get_json()['normalization_review']['has_overrides'] is override


@pytest.mark.parametrize('mutation',['reviewer','boolean','unknown','duplicate','forged_source','forged_ai','unknown_key'])
def test_invalid_or_forged_decisions_are_rejected(mutation):
    doc,source=fixture_pair()
    d=decision(action='manual',value=24)
    if mutation=='reviewer': d['reviewer']=''
    if mutation=='boolean': d['value']=True
    if mutation=='unknown': d['item_id']='not-an-item'
    if mutation=='forged_source': d.update(action='source',value=123)
    if mutation=='forged_ai': d.update(action='ai',value=123)
    if mutation=='unknown_key': d['approved']=True
    saved=envelope(doc,source,[d,d] if mutation=='duplicate' else [d])
    assert review(saved,source).status_code==400


def test_save_reload_and_changed_source_detection():
    doc,source=fixture_pair()
    saved=envelope(doc,source,[decision(action='manual',value=24)])
    first=review(saved,source).get_json()
    assert 'reviewed_document' in first
    restored=review(first['reviewed_document'],source).get_json()
    assert restored['items'][0]['quantity']==24
    assert restored['decision_audit']==first['decision_audit']
    assert review(first['reviewed_document'],change_source(source,'D3',13)).status_code==400


def test_stale_amount_blocks_cash_until_explicitly_resolved():
    doc,source=fixture_pair()
    saved=envelope(doc,source,[decision(action='manual',value=24)])
    result=review(saved,source).get_json()
    assert any(c['code']=='EFFECTIVE_AMOUNT_CONFLICT' for c in result['calculation_checks'])
    client=create_app().test_client()
    assert client.post('/api/normalization/analyze',data=analysis_form(saved,source,cashflow_mode='compare')).status_code==422
    saved['decisions'].append(decision(field='amount',action='manual',value=240))
    response=client.post('/api/normalization/analyze',data=analysis_form(saved,source,cashflow_mode='compare'))
    assert response.status_code==200,response.get_json()
    assert response.get_json()['metrics']['boq_total']==340


def test_unit_change_requires_explicit_quantity_review_and_remains_compatible():
    doc,source=fixture_pair()
    saved=envelope(doc,source,[decision(field='unit',action='manual',value='kg')])
    result=review(saved,source).get_json()
    assert any(c['code']=='UNIT_QUANTITY_REVIEW' for c in result['calculation_checks'])
    assert any(c['code']=='UNSUPPORTED_OPERATION_UNIT' for c in result['calculation_checks'])
    assert result['admission_ready'] is False


@pytest.mark.parametrize('problem',['missing_row','bad_ref','unsupported'])
def test_field_override_cannot_waive_unrelated_blockers(problem):
    doc,source=fixture_pair()
    if problem=='missing_row':
        doc['items'].pop()
        doc['extraction_summary'].update(emitted_item_count=1,source_leaf_item_count=1)
    elif problem=='bad_ref': doc['items'][0]['quantity_ref']='E3'
    else: doc['items'][0].update(operation='filter_placement',work_package='filter_media')
    saved=envelope(doc,source,[decision(action='manual',value=24)])
    result=review(saved,source).get_json()
    assert result['admission_ready'] is False
    response=create_app().test_client().post('/api/normalization/analyze',data=analysis_form(saved,source))
    assert response.status_code==422


def test_evidence_correction_uses_only_server_derived_patch():
    doc,source=fixture_pair()
    doc['items'][0]['raw_description']='Fabricated quote'
    doc['items'][0]['quantity_ref']='E3'
    initial=review(doc,source).get_json()
    repair=next(o for o in initial['resolution_options'] if o['item_id']=='I3' and o['field']=='source_evidence')
    saved=envelope(doc,source,[decision(field='source_evidence',action='source',value=repair['source_value'])])
    result=review(saved,source).get_json()
    assert result['admission_ready'] is True
    assert result['document']['items'][0]['raw_description']=='Fabricated quote'
    assert result['items'][0]['raw_description']=='Filling available excavated earth'
    saved['decisions'][0]['value']['quantity_ref']='F3'
    assert review(saved,source).status_code==400


def test_missing_prices_can_be_entered_without_pretending_they_are_source_prices():
    doc,source=fixture_pair(False)
    saved=envelope(doc,source,[decision(field='rate',action='manual',value=10),decision(field='amount',action='manual',value=120)])
    result=review(saved,source).get_json()
    assert result['items'][0]['amount']==120
    assert result['items'][1]['amount'] is None
    assert result['priced_total'] is None
    assert result['has_overrides'] is True
    saved['decisions'][1]['value']=0
    result=review(saved,source).get_json()
    assert result['items'][0]['amount']==0
    assert result['unpriced_item_count']==1


def test_saved_approval_is_not_authority_and_overrides_survive_export_and_replan():
    doc,source=fixture_pair()
    saved=envelope(doc,source,[decision(action='manual',value=24)])
    client=create_app().test_client()
    assert client.post('/api/normalization/analyze',data=analysis_form(saved,source,review_confirmed='false')).status_code==422
    response=client.post('/api/normalization/analyze',data=analysis_form(saved,source))
    assert response.status_code==200,response.get_json()
    result=response.get_json()
    assert any(f['code']=='PLANNER_INPUT_OVERRIDES' for f in result['findings'])
    replan=client.post(f'/api/jobs/{result["job_id"]}/replan',json={'overrides':[{'id':'MOB','duration_days':3}]}).get_json()
    assert replan['normalization_review']['has_overrides'] is True
    assert any(f['code']=='PLANNER_INPUT_OVERRIDES' for f in replan['findings'])
    book=load_workbook(io.BytesIO(client.get(f'/api/jobs/{result["job_id"]}/export.xlsx').data))
    assert 'Input Decisions' in book.sheetnames
    assert book['Priced BOQ']['D5'].value==24
    values=[v for row in book['Input Decisions'].values for v in row]
    assert 12 in values and 24 in values and 'Verified revised site scope' in values


def test_scope_confirmation_still_resolves_the_existing_model_scope_question():
    doc,source=fixture_pair()
    doc['issues']=[dict(id='SCOPE',code='QUANTITY_SCOPE_UNRESOLVED',item_ids=['I3'],source_refs=[],
        message='Confirm total quantities',blocks=['schedule'],question_for_reviewer='Project totals?')]
    doc['extraction_summary']['issue_count']=1
    response=create_app().test_client().post('/api/normalization/analyze',data=analysis_form(envelope(doc,source),source))
    assert response.status_code==200,response.get_json()


def test_unit_change_accepts_explicit_quantity_review_without_reason():
    doc,source=fixture_pair()
    doc['items'][0].update(operation='reinforcement_fixing',work_package='reinforcement')
    saved=envelope(doc,source,[decision(field='unit',action='manual',value='kg'),
        decision(action='source',value=12,reason='')])
    result=review(saved,source).get_json()
    assert not any(c['code']=='UNIT_QUANTITY_REVIEW' for c in result['calculation_checks'])


@pytest.mark.parametrize('action',['ai','manual'])
@pytest.mark.parametrize('reason',['', ' ', None])
def test_reason_is_optional_in_saved_decisions_and_explicitly_recorded(action,reason):
    doc,source=fixture_pair()
    doc['items'][0]['quantity']=24
    d=decision(value=24,action=action,reason=reason)
    if reason is None: del d['reason']
    result=review(envelope(doc,source,[d]),source)
    assert result.status_code==200,result.get_json()
    assert result.get_json()['decision_audit'][0]['reason']=='Not provided'
    assert result.get_json()['has_overrides'] is True


def test_bulk_ai_decisions_do_not_waive_unsupported_engineering():
    doc,source=fixture_pair()
    doc['items'][0].update(operation='filter_placement',work_package='filter_media')
    initial=review(doc,source).get_json()
    choices=[decision(field=o['field'],value=o['proposed_value'],action='ai',item=o['item_id'],reason='')
             for o in initial['resolution_options'] if 'ai' in o['actions']]
    saved=envelope(doc,source,choices)
    result=review(saved,source).get_json()
    assert any(c['code']=='UNSUPPORTED_OPERATION_UNIT' for c in result['calculation_checks'])
    assert create_app().test_client().post('/api/normalization/analyze',data=analysis_form(saved,source)).status_code==422


def test_source_description_cannot_be_waived_as_a_planning_description():
    doc,source=fixture_pair()
    doc['items'][0]['raw_description']='invented'
    saved=envelope(doc,source,[decision(field='normalized_description',action='manual',value='Filling soil')])
    result=review(saved,source).get_json()
    assert result['admission_ready'] is False
    assert any(c['code']=='SOURCE_DESCRIPTION_MISMATCH' for c in result['calculation_checks'])


def test_save_endpoint_produces_reloadable_draft_and_does_not_save_approval_flags():
    doc,source=fixture_pair()
    saved=envelope(doc,source,[decision(action='manual',value=24)])
    client=create_app().test_client()
    response=client.post('/api/normalization/save',data=analysis_form(saved,source))
    assert response.status_code==200,response.get_json()
    result=response.get_json()
    assert result['review_state']=='draft'
    assert 'review_confirmed' not in result['review_settings']
    assert review(result,source).get_json()['items'][0]['quantity']==24
